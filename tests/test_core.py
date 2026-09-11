import asyncio
import io
import json
import os
import sys
import time
import zipfile
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from pypdf import PdfWriter
from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject

from clara.config import Config
from clara.store import Store
from clara.app import create_app
from clara.agent import AgentManager
from clara.auth import clean_environment
from clara.skills import import_skill, list_skills
from clara.tools import permitted_path, search_files, read_document, run_command, publish_artifact


@pytest.fixture
def config(tmp_path):
    cfg = Config(tmp_path / "data")
    cfg.initialize()
    return cfg


def make_pdf(path, text="Client Rohit Sharma - engagement letter - TEST ONLY"):
    writer = PdfWriter()
    page = writer.add_blank_page(width=612, height=792)
    font = DictionaryObject({NameObject("/Type"): NameObject("/Font"), NameObject("/Subtype"): NameObject("/Type1"), NameObject("/BaseFont"): NameObject("/Helvetica")})
    page[NameObject("/Resources")] = DictionaryObject({NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)})})
    stream = DecodedStreamObject()
    stream.set_data(f"BT /F1 12 Tf 40 740 Td ({text}) Tj ET".encode())
    page[NameObject("/Contents")] = writer._add_object(stream)
    writer.write(path)


def test_search_moved_pdf_and_verify_content(config):
    target = config.workspace / "moved" / "2025" / "Sharma-Engagement.PDF"
    target.parent.mkdir(parents=True)
    make_pdf(target)
    result = search_files(config, ".", "*sharma*.pdf")
    assert len(result["matches"]) == 1
    assert result["matches"][0]["path"] == str(target)
    doc = read_document(config, str(target))
    assert doc["total_pages"] == 1
    assert "Rohit Sharma" in doc["pages"][0]["text"]


def test_file_tools_block_traversal_symlinks_and_skill_writes(config, tmp_path):
    secret = tmp_path / "outside.txt"
    secret.write_text("outside")
    with pytest.raises(ValueError):
        permitted_path(config, str(secret))
    with pytest.raises(ValueError):
        permitted_path(config, "../../outside.txt")
    (config.workspace / "link.txt").symlink_to(secret)
    with pytest.raises(ValueError):
        permitted_path(config, "link.txt")
    with pytest.raises(ValueError):
        permitted_path(config, ".claude/skills/evil/SKILL.md", write=True)
    assert search_files(config, ".", "link.txt")["matches"] == []


def test_artifact_is_an_immutable_copy(config):
    store = Store(config.data / "test.sqlite")
    c = store.create_conversation()
    job = store.create_job(c["id"], "test", "ask", [])
    source = config.workspace / "test.txt"
    source.write_text("original result")
    result = publish_artifact(config, store, job, "test.txt")
    source.write_text("changed later")
    record = store.one("SELECT * FROM files WHERE id=?", (result["id"],))
    assert Path(record["path"]).read_text() == "original result"
    assert result["size"] == 15


def test_read_roots_allow_reads_but_not_writes(config, tmp_path):
    directory = tmp_path / "client-files"
    directory.mkdir()
    file = directory / "hello.txt"
    file.write_text("hello")
    settings = config.settings()
    settings["read_roots"] = [str(directory)]
    config.save_settings(settings)
    assert read_document(config, str(file))["text"] == "hello"
    with pytest.raises(ValueError):
        permitted_path(config, str(file), write=True)


def test_utf16_meeting_transcripts_are_readable(config):
    file = config.workspace / 'meeting.txt'
    file.write_text('Synthetic meeting transcript: use a working copy.', encoding='utf-16')
    assert read_document(config, 'meeting.txt')['text']=='Synthetic meeting transcript: use a working copy.'


def skill_zip(files):
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as archive:
        for name, content in files.items():
            archive.writestr(name, content)
    return buf.getvalue()


SKILL = "---\nname: test-procedure\ndescription: A test procedure for document finding.\n---\nVerify the exact client.\n"


def test_skill_import_keeps_relative_resources(config):
    data = skill_zip({"procedure/SKILL.md": SKILL, "procedure/references/example.txt": "example"})
    import_skill(config, "procedure.zip", data)
    assert (config.skills / "test-procedure/references/example.txt").read_text() == "example"
    assert "test-procedure" in [s["name"] for s in list_skills(config)]
    with pytest.raises(ValueError, match="exists"):
        import_skill(config, "procedure.zip", data)


@pytest.mark.parametrize("bad_path", ["../escape.txt", "/absolute.txt", "C:\\escape.txt", "procedure/../../escape.txt", "procedure/.claude/settings.json"])
def test_skill_zip_rejects_unsafe_paths(config, bad_path):
    with pytest.raises(ValueError):
        import_skill(config, "bad.zip", skill_zip({"procedure/SKILL.md": SKILL, bad_path: "bad"}))
    assert not (config.skills / "test-procedure").exists()


def test_skill_zip_rejects_symlink(config):
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w") as archive:
        archive.writestr("procedure/SKILL.md", SKILL)
        entry = zipfile.ZipInfo("procedure/link")
        entry.create_system = 3
        entry.external_attr = 0o120777 << 16
        archive.writestr(entry, "/etc/passwd")
    with pytest.raises(ValueError):
        import_skill(config, "bad.zip", output.getvalue())


def test_commands_use_venv_and_have_timeout(config):
    async def scenario():
        result = await run_command(config, 'python -c "import sys; print(sys.executable)"')
        assert result["exit_code"] == 0
        assert str(Path(sys.executable).parent) in result["output"]
        with pytest.raises(TimeoutError):
            await run_command(config, 'python -c "import time; time.sleep(30)"', timeout_seconds=1)
    asyncio.run(scenario())


def test_restart_marks_jobs_interrupted_without_replay(config):
    store = Store(config.data / "test.sqlite")
    c = store.create_conversation()
    job = store.create_job(c["id"], "Do a task", "ask", [])
    store.status(job["id"], "running")
    store.recover()
    assert store.job(job["id"])["status"] == "interrupted"
    assert store.events(c["id"])[-1]["data"]["status"] == "interrupted"


def test_billing_environment_is_removed_not_forwarded():
    result = clean_environment({"PATH": "/bin", "HOME": "/home/person", "ANTHROPIC_API_KEY": "secret",
        "ANTHROPIC_BASE_URL": "https://other.example", "CLAUDE_CODE_USE_BEDROCK": "1",
        "CLAUDE_CODE_OAUTH_TOKEN": "token", "AWS_SECRET_ACCESS_KEY": "secret", "CLAUDE_CONFIG_DIR": "/other"})
    assert result == {"PATH": "/bin", "HOME": "/home/person"}


def test_office_files_created_through_real_command_tool(config):
    script = config.workspace / 'office-check.py'
    script.write_text("""from docx import Document
from pptx import Presentation
from openpyxl import Workbook
document=Document();document.add_heading('Synthetic Clara test',0);document.save('outputs/test.docx')
slides=Presentation();slides.slides.add_slide(slides.slide_layouts[0]);slides.save('outputs/test.pptx')
book=Workbook();book.active['A1']='Synthetic test';book.save('outputs/test.xlsx')
""")
    result = asyncio.run(run_command(config, 'python office-check.py'))
    assert result['exit_code']==0, result
    from docx import Document
    from pptx import Presentation
    from openpyxl import load_workbook
    assert Document(config.workspace/'outputs/test.docx').paragraphs[0].text=='Synthetic Clara test'
    assert len(Presentation(config.workspace/'outputs/test.pptx').slides)==1
    assert load_workbook(config.workspace/'outputs/test.xlsx').active['A1'].value=='Synthetic test'


class NoInferenceManager(AgentManager):
    """Test executor; never selectable by the production app or UI."""
    async def execute(self, job):
        self.store.status(job["id"], "running")
        await asyncio.sleep(0.15)
        self.store.event(job["conversation_id"], job["id"], "assistant", {"text": "Offline test event"})
        self.store.status(job["id"], "completed")


@pytest.fixture
def client(config, monkeypatch):
    async def auth(*args):
        return {"connected": False, "message": "Not signed in", "api_fallback": False}
    monkeypatch.setattr("clara.app.auth_status", auth)
    app = create_app(config, manager_factory=NoInferenceManager, access_token="test-local-access")
    with TestClient(app, base_url="http://127.0.0.1:8876") as client:
        client.headers["x-clara-request"] = "1"
        client.app_instance = app
        yield client


def login(client):
    assert client.post("/api/session", json={"token": "test-local-access"}).status_code == 200


def test_loopback_auth_origin_and_content_security(client):
    assert client.get("/").status_code == 200
    assert "frame-ancestors 'none'" in client.get("/").headers["content-security-policy"]
    assert client.get("/api/status").status_code == 401
    assert client.get("/api/status", headers={"host": "attacker.example:8876"}).status_code == 403
    login(client)
    assert client.get("/api/status").status_code == 200
    assert client.post("/api/conversations", json={}, headers={"origin": "https://attacker.example"}).status_code == 403
    assert client.post("/api/conversations", json={}, headers={"x-clara-request": ""}).status_code == 403


def test_api_upload_task_and_event_persistence(client, config):
    login(client)
    cid = client.post("/api/conversations", json={}).json()["id"]
    upload = client.post(f"/api/conversations/{cid}/upload", files={"file": ("../../test.txt", b"test data")})
    assert upload.status_code == 200
    file = upload.json()
    assert file["name"] == "test.txt"
    assert client.get("/api/files/" + file["id"]).content == b"test data"
    job = client.post(f"/api/conversations/{cid}/messages", json={"text": "Do the offline test", "attachments": [file["id"]]}).json()
    assert "id" in job
    deadline = time.monotonic() + 3
    while time.monotonic() < deadline:
        jobs = client.get(f"/api/conversations/{cid}").json()["jobs"]
        if jobs[0]["status"] == "completed":
            break
        time.sleep(0.05)
    assert jobs[0]["status"] == "completed"
    store = client.app_instance.state.store
    events = store.events(cid)
    assert any(e["kind"] == "assistant" for e in events)
    assert store.events(cid, events[-2]["id"]) == events[-1:]
    cid2 = client.post("/api/conversations", json={}).json()["id"]
    assert client.post(f"/api/conversations/{cid2}/messages", json={"text": "x", "attachments": [file["id"]]}).status_code == 400


def test_api_skill_create_import_and_settings_validation(client, config):
    login(client)
    assert client.put("/api/skills/test-procedure", json={"content": SKILL}).status_code == 200
    assert client.get("/api/skills").json()[-1]["content"]
    settings = client.get("/api/settings").json()
    settings["read_roots"] = [str(config.data / "does-not-exist")]
    assert client.put("/api/settings", json=settings).status_code == 400
    settings["read_roots"] = []
    settings["model"] = "sonnet"
    assert client.put("/api/settings", json=settings).status_code == 200


def test_queue_serializes_cancels_and_resolves_questions(config):
    class WaitingManager(AgentManager):
        def __init__(self, *args):
            super().__init__(*args)
            self.started = []
        async def execute(self, job):
            self.started.append(job["id"])
            self.store.status(job["id"], "running")
            await self.request_input(job, "question", {"question": "Which client?"})
            self.store.status(job["id"], "completed")
    async def scenario():
        store = Store(config.data / "queue.sqlite")
        manager = WaitingManager(config, store)
        await manager.start()
        one, two = store.create_conversation(), store.create_conversation()
        first = manager.submit(one["id"], "one", "ask", [])
        second = manager.submit(two["id"], "two", "ask", [])
        for _ in range(100):
            if manager.pending:
                break
            await asyncio.sleep(0.01)
        assert manager.started == [first["id"]]
        with pytest.raises(ValueError):
            manager.submit(one["id"], "duplicate", "ask", [])
        await manager.cancel(first["id"])
        for _ in range(100):
            if manager.active_job == second["id"] and manager.pending:
                break
            await asyncio.sleep(0.01)
        assert store.job(first["id"])["status"] == "cancelled"
        manager.answer(next(iter(manager.pending)), "Test client")
        await asyncio.wait_for(manager.queue.join(), 2)
        assert store.job(second["id"])["status"] == "completed"
        assert not manager.pending
        await manager.close()
    asyncio.run(scenario())
