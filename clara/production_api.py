"""Operator-only management endpoints, protected by the dashboard session middleware."""
import asyncio
import json
from fastapi import HTTPException,UploadFile,File,Form
from pydantic import BaseModel,Field
from .knowledge import Knowledge,Memory,LIBRARIES,decode
from .workflows import Workflows,TEMPLATES


class KnowledgeInput(BaseModel):
    title:str=Field(min_length=1,max_length=180)
    text:str=Field(min_length=1,max_length=500000)
    library:str
    metadata:dict[str,str]
    client_key:str=''

class ReviewInput(BaseModel):
    action:str='approve'
    note:str=Field(min_length=1,max_length=4000)
    share:bool=False

class WorkflowInput(BaseModel):
    kind:str
    context:dict

class BudgetInput(BaseModel):
    max_turns:int|None=Field(ge=1)
    task_timeout_minutes:int=Field(ge=1,le=480)
    max_budget_usd:float|None=Field(default=None,gt=0,le=1000,allow_inf_nan=False)
    note:str=Field(min_length=1,max_length=4000)

class UsageReviewInput(BaseModel):
    job_fingerprints:dict[str,str]
    note:str=Field(min_length=1,max_length=4000)

def install_routes(app,config,store,manager):
    knowledge=Knowledge(store);memory=Memory(store);wf=Workflows(store,config)
    def idle():
        if manager.active_job or not manager.queue.empty():
            raise ValueError('Finish or stop active tasks before changing reviewed knowledge, memory or workflow budgets.')
    def conversation(cid):
        if not store.conversation(cid): raise HTTPException(404,'Conversation not found.')
    @app.get('/api/knowledge')
    async def knowledge_list(client_key:str=''):
        return {'libraries':LIBRARIES,'records':knowledge.list('client:'+client_key if client_key else 'firm')}
    @app.post('/api/knowledge')
    async def knowledge_add(body:KnowledgeInput):
        idle()
        return await asyncio.to_thread(knowledge.add,body.title,body.text,body.library,body.metadata,'client:'+body.client_key if body.client_key else 'firm')
    @app.post('/api/knowledge/import')
    async def knowledge_import(file:UploadFile=File(...),details:str=Form(...)):
        idle()
        from pathlib import Path
        import io
        data=await file.read(10_000_001)
        if len(data)>10_000_000: raise ValueError('Choose a reference under 10 MB.')
        meta=json.loads(details);suffix=Path(file.filename or '').suffix.lower()
        def extract():
            if suffix=='.pdf':
                from pypdf import PdfReader
                reader=PdfReader(io.BytesIO(data))
                if reader.is_encrypted or len(reader.pages)>200:
                    raise ValueError('Use a readable PDF of at most 200 pages, or split the manual into sections.')
                return '\n\n'.join(f'Page {i+1}\n'+(p.extract_text() or '') for i,p in enumerate(reader.pages))
            if suffix=='.docx':
                from docx import Document
                doc=Document(io.BytesIO(data))
                return '\n\n'.join([p.text for p in doc.paragraphs]+[' | '.join(c.text for c in row.cells) for t in doc.tables for row in t.rows])
            if suffix not in {'.txt','.md'}: raise ValueError('Use PDF, DOCX, text or Markdown.')
            return data.decode('utf-16' if data.startswith((b'\xff\xfe',b'\xfe\xff')) else 'utf-8-sig')
        text=await asyncio.to_thread(extract)
        if suffix=='.pdf' and not any(line.strip() and not line.startswith('Page ') for line in text.splitlines()):
            raise ValueError('No text found in this scanned PDF. Import an OCR text version with the source reference.')
        body=KnowledgeInput(**{**meta,'text':text})
        return await asyncio.to_thread(knowledge.add,body.title,body.text,body.library,body.metadata,'client:'+body.client_key if body.client_key else 'firm')
    @app.post('/api/knowledge/{kid}/retire')
    async def knowledge_retire(kid:str):
        idle();knowledge.retire(kid);return {'retired':True}
    @app.get('/api/memories')
    async def memories(): return memory.list()
    @app.post('/api/memories/{mid}/review')
    async def review_memory(mid:str,body:ReviewInput):
        idle();return memory.review(mid,body.action,body.note,body.share)
    @app.get('/api/conversations/{cid}/workflow')
    async def workflow(cid:str):
        conversation(cid)
        jobs=store.rows('SELECT * FROM jobs WHERE conversation_id=? ORDER BY created DESC LIMIT 1',(cid,))
        from .operations import Operations
        return {'workflow':wf.snapshot(cid),'templates':TEMPLATES,
                'operations':Operations(store).list(jobs[0]) if jobs else [],
                'evidence':[decode(r,'payload') for r in store.rows('SELECT e.* FROM evidence e JOIN jobs j ON j.id=e.job_id WHERE j.conversation_id=? AND e.kind!=? ORDER BY e.created DESC LIMIT 100',(cid,'tool_observation'))]}
    @app.post('/api/operations/{oid}/confirm-absence')
    async def operation_absence(oid:str,body:ReviewInput):
        idle()
        from .operations import Operations
        return Operations(store).confirm_absence(oid,body.note)
    @app.post('/api/conversations/{cid}/workflow')
    async def begin_workflow(cid:str,body:WorkflowInput):
        idle();conversation(cid)
        return wf.begin({'id':None,'conversation_id':cid},body.kind,body.context)
    @app.post('/api/conversations/{cid}/workflow/review')
    async def review_workflow(cid:str,body:ReviewInput):
        idle();conversation(cid);return await asyncio.to_thread(wf.review,cid,body.note)
    @app.put('/api/conversations/{cid}/workflow/budget')
    async def workflow_budget(cid:str,body:BudgetInput):
        idle();conversation(cid)
        w=wf.get(cid)
        if not w: raise ValueError('No workflow exists in this conversation.')
        limits=body.model_dump(exclude={'note'})
        store.execute('UPDATE workflows SET limits=? WHERE id=?',(json.dumps(limits),w['id']))
        store.event(cid,None,'workflow_budget',{'limits':limits,'note':body.note})
        return wf.snapshot(cid)
    @app.post('/api/conversations/{cid}/workflow/review-usage')
    async def review_usage(cid:str,body:UsageReviewInput):
        idle();conversation(cid)
        return wf.review_usage(cid,body.job_fingerprints,body.note)
