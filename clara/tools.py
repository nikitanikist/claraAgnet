"""General capabilities, independent of any tax workflow or screen layout."""
import asyncio
import fnmatch
import json
import os
import shutil
import signal
import sys
import tempfile
import time
from pathlib import Path

from claude_agent_sdk import create_sdk_mcp_server, tool
from .auth import clean_environment
from .store import new_id


def inside(path, roots):
    return any(path.is_relative_to(root) for root in roots)


def permitted_path(config, value, write=False):
    path = Path(value).expanduser()
    if not path.is_absolute():
        path = config.workspace / path
    path = path.resolve()
    roots = [config.workspace.resolve()] if write else config.read_roots()
    if not inside(path, roots):
        raise ValueError("Path is outside the configured folders. Add its folder in Settings first.")
    if write and (path.is_relative_to(config.workspace / ".claude") or path == config.workspace):
        raise ValueError("Use the Skills editor to change Clara's instructions.")
    return path


def search_files(config, root, pattern, max_results=100):
    base = permitted_path(config, root)
    if not base.is_dir():
        raise ValueError("Search root is not a directory.")
    results, skipped, scanned = [], 0, 0
    deadline = time.monotonic() + 15
    for directory, folders, files in os.walk(base, followlinks=False):
        folders[:] = [f for f in folders if f not in {"node_modules", ".git", ".venv", "$RECYCLE.BIN"}
                      and not (Path(directory) / f).is_symlink()]
        for filename in files:
            scanned += 1
            if time.monotonic() > deadline or scanned > 200_000:
                return {"matches": results, "truncated": True, "scanned": scanned, "skipped": skipped}
            if fnmatch.fnmatch(filename.casefold(), pattern.casefold()):
                path = (Path(directory) / filename).resolve()
                if not inside(path, config.read_roots()):
                    skipped += 1
                    continue
                try:
                    results.append({"path": str(path), "size": path.stat().st_size})
                except OSError:
                    skipped += 1
                if len(results) >= max_results:
                    return {"matches": results, "truncated": True, "scanned": scanned, "skipped": skipped}
    return {"matches": results, "truncated": False, "scanned": scanned, "skipped": skipped}


def read_document(config, value, start_page=1, page_count=10):
    path = permitted_path(config, value)
    if path.stat().st_size > 50_000_000:
        raise ValueError("Document exceeds 50 MB. Narrow or split the file first.")
    start_page = max(1, start_page)
    page_count = min(30, max(1, page_count))
    if path.suffix.lower() == ".pdf":
        from pypdf import PdfReader
        reader = PdfReader(path)
        pages = [{"page": i + 1, "text": (reader.pages[i].extract_text() or "")[:15000]}
                 for i in range(start_page - 1, min(len(reader.pages), start_page - 1 + page_count))]
        return {"path": str(path), "total_pages": len(reader.pages), "pages": pages,
                "note": "Empty text may mean scanned pages; OCR or visual inspection is then required."}
    if path.suffix.lower() == ".docx":
        from docx import Document
        doc = Document(path)
        text = "\n".join(p.text for p in doc.paragraphs)
        text += "\n" + "\n".join(" | ".join(c.text for c in row.cells) for t in doc.tables for row in t.rows)
    else:
        if path.suffix.lower() in {".xlsx", ".xls", ".pptx", ".ppt", ".zip", ".exe"}:
            raise ValueError("This is not a text/PDF/DOCX file. For Office files use the office-documents skill and the installed Python library.")
        raw = path.read_bytes()[:200_000]
        encoding = "utf-16" if raw.startswith((b"\xff\xfe", b"\xfe\xff")) else "utf-8-sig"
        text = raw.decode(encoding, errors="replace")
    return {"path": str(path), "text": text[:100_000], "truncated": len(text) > 100_000}


async def terminate_tree(process):
    if process.returncode is not None:
        return
    if os.name == "nt":
        killer = await asyncio.create_subprocess_exec("taskkill", "/PID", str(process.pid), "/T", "/F",
            stdout=asyncio.subprocess.DEVNULL, stderr=asyncio.subprocess.DEVNULL)
        await killer.wait()
    else:
        try:
            os.killpg(process.pid, signal.SIGKILL)
        except ProcessLookupError:
            pass
    await process.wait()


async def run_command(config, command, cwd=".", timeout_seconds=60):
    directory = permitted_path(config, cwd)
    if not directory.is_relative_to(config.workspace) or directory.is_relative_to(config.workspace / ".claude"):
        raise ValueError("Run commands from Clara's workspace, outside its instruction directory.")
    timeout_seconds = min(300, max(1, timeout_seconds))
    shell = ["powershell.exe", "-NoProfile", "-NonInteractive", "-Command", command] if os.name == "nt" else ["/bin/bash", "-c", command]
    environment = clean_environment()
    environment["PATH"] = str(Path(sys.executable).parent) + os.pathsep + environment.get("PATH", "")
    environment["CLARA_WORKSPACE"] = str(config.workspace)
    with tempfile.TemporaryFile() as output:
        process = await asyncio.create_subprocess_exec(*shell, cwd=directory, env=environment,
            stdout=output, stderr=asyncio.subprocess.STDOUT, start_new_session=os.name != "nt")
        try:
            await asyncio.wait_for(process.wait(), timeout_seconds)
        except (TimeoutError, asyncio.CancelledError):
            await terminate_tree(process)
            raise
        output.seek(0, 2)
        length = output.tell()
        output.seek(0)
        text = output.read(48_000).decode("utf-8", errors="replace")
        return {"exit_code": process.returncode, "output": text, "truncated": length > 48_000}


def publish_artifact(config, store, job, source, title=""):
    path = permitted_path(config, source)
    if not path.is_file() or path.stat().st_size > 100_000_000:
        raise ValueError("Choose a file of 100 MB or less.")
    fid = new_id()
    name = path.name
    target = config.data / "artifacts" / (fid + path.suffix)
    shutil.copyfile(path, target)
    store.execute("INSERT INTO files VALUES(?,?,?,?,?,?,?,?)", (fid, job["conversation_id"], job["id"],
                  "artifact", name, str(target), target.stat().st_size, time.time()))
    # Kept as a content snapshot, so later source edits cannot silently alter an attachment.
    result = {"id": fid, "name": name, "title": title or name, "size": target.stat().st_size,
              "url": f"/api/files/{fid}"}
    store.event(job["conversation_id"], job["id"], "artifact", result)
    return result


def tool_server(config, store, job, request_input):
    def wrap(fn):
        async def call(args):
            try:
                result = await fn(args)
                return {"content": [{"type": "text", "text": json.dumps(result, ensure_ascii=False)}]}
            except (ValueError, OSError, TimeoutError) as error:
                return {"isError": True, "content": [{"type": "text", "text": str(error)}]}
        return call

    async def search(args):
        return await asyncio.to_thread(search_files, config, args["root"], args["pattern"])

    async def read(args):
        return await asyncio.to_thread(read_document, config, args["path"], args.get("start_page", 1), args.get("page_count", 10))

    async def write(args):
        path = permitted_path(config, args["path"], write=True)
        if len(args["content"].encode()) > 2_000_000:
            raise ValueError("Text file exceeds 2 MB.")
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text(args["content"], encoding="utf-8")
        return {"path": str(path), "bytes": path.stat().st_size}

    async def command(args):
        if job["mode"] != "autonomous":
            approved = await request_input(job, "approval", {"tool": "run_command", "input": args})
            if approved != "allow":
                return {"denied": True, "message": "The user declined this command."}
        # OS commands have account-level access. Folder restrictions on file tools
        # are not an OS sandbox; the dashboard makes this distinction explicit.
        return await run_command(config, args["command"], args.get("cwd", "."), args.get("timeout_seconds", 60))

    async def artifact(args):
        return publish_artifact(config, store, job, args["path"], args.get("title", ""))

    async def ask(args):
        return {"answer": await request_input(job, "question", {"question": args["question"]})}

    async def environment(args):
        from .connectors import connector_status
        return {"platform": sys.platform, "python": sys.executable, "workspace": str(config.workspace),
                "read_roots": [str(p) for p in config.read_roots()], "mode": job["mode"],
                "libraries": ["pypdf", "docx", "pptx", "openpyxl"],
                "browser_enabled": config.settings()["browser_enabled"],
                "desktop_enabled": config.settings()["desktop_enabled"] and os.name == "nt",
                "connectors":connector_status(config)}

    definitions = [
        ("environment", "Get platform, folders, Python executable and enabled capabilities.", {}, environment),
        ("search_files", "Search file names recursively, case-insensitively. Use wildcards; inspect candidates before returning a client document.", {"root": str, "pattern": str}, search),
        ("read_document", "Read text or DOCX, or extract numbered PDF pages. Scanned PDF images need separate visual inspection/OCR.", {"type": "object", "properties": {"path": {"type": "string"}, "start_page": {"type": "integer"}, "page_count": {"type": "integer"}}, "required": ["path"]}, read),
        ("write_text", "Write a UTF-8 file inside Clara's workspace. Also useful to prepare a Python script for execution.", {"path": str, "content": str}, write),
        ("run_command", "Run a command in PowerShell on Windows or Bash on Mac. Python and document libraries are installed. Commands have the OS user's access, not a sandbox. Ask mode requires confirmation.", {"type": "object", "properties": {"command": {"type": "string"}, "cwd": {"type": "string"}, "timeout_seconds": {"type": "integer"}}, "required": ["command"]}, command),
        ("publish_artifact", "Attach a verified, real file to the conversation for the user to download. Does not send anything to an external recipient.", {"type": "object", "properties": {"path": {"type": "string"}, "title": {"type": "string"}}, "required": ["path"]}, artifact),
        ("ask_user", "Ask a necessary clarification and wait for the user's answer in the dashboard.", {"question": str}, ask),
    ]
    from .production_tools import definitions as production_definitions
    definitions += production_definitions(config,store,job)
    return create_sdk_mcp_server(name="clara", version="0.2.2", tools=[tool(name, desc, schema)(wrap(fn)) for name, desc, schema, fn in definitions])
